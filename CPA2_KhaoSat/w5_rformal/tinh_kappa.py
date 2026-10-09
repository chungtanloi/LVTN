"""
tinh_kappa.py — tính mức đồng thuận giữa hai người gán nhãn trên danh mục R-Formal.

Chạy:  python3 tinh_kappa.py phieu_da_dien.xlsx

Đọc phiếu đã điền (sinh bởi make_review_kit.py) và đối chiếu với r_formal_v1.json.
Báo cáo, cho từng trục (lop F/P/N và du_lieu co/khong):
  · tỷ lệ đồng thuận thô
  · Cohen's kappa, kèm cách đọc theo thang Landis & Koch
  · ma trận nhầm lẫn
  · danh sách đơn vị bất đồng, để hai người ngồi lại giải quyết

Không phụ thuộc scikit-learn — kappa được tính trực tiếp, vài dòng số học.
"""
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
CAT = HERE / "r_formal_v1.json"


def kappa(a, b, nhan):
    """Cohen's kappa cho hai danh sách nhãn song song."""
    n = len(a)
    if n == 0:
        return float("nan")
    po = sum(1 for x, y in zip(a, b) if x == y) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum((ca[k] / n) * (cb[k] / n) for k in nhan)
    if pe == 1:
        return float("nan")
    return (po - pe) / (1 - pe)


def doc_kappa(k):
    if k != k:
        return "không xác định"
    if k < 0:    return "tệ hơn ngẫu nhiên"
    if k < 0.20: return "rất thấp"
    if k < 0.40: return "thấp"
    if k < 0.60: return "trung bình"
    if k < 0.80: return "tốt (substantial)"
    return "rất tốt (almost perfect)"


def bao_cao(ten_truc, cap, nhan, ids):
    a = [x for x, _ in cap]
    b = [y for _, y in cap]
    n = len(cap)
    print("\n" + "=" * 74)
    print(f"TRỤC: {ten_truc}   (n = {n})")
    print("=" * 74)
    if n == 0:
        print("  Không có ô nào được điền — bỏ qua.")
        return
    dong_y = sum(1 for x, y in zip(a, b) if x == y)
    k = kappa(a, b, nhan)
    print(f"  Đồng thuận thô : {dong_y}/{n} = {100*dong_y/n:.1f}%")
    print(f"  Cohen's kappa  : {k:.3f}  → {doc_kappa(k)}")

    print("\n  Ma trận nhầm lẫn (hàng = người 1 / danh mục, cột = người 2):")
    m = defaultdict(int)
    for x, y in cap:
        m[(x, y)] += 1
    print("        " + "".join(f"{c:>8}" for c in nhan))
    for r in nhan:
        print(f"    {r:<4}" + "".join(f"{m[(r,c)]:>8}" for c in nhan))

    lech = [(i, x, y) for i, (x, y) in zip(ids, cap) if x != y]
    if lech:
        print(f"\n  {len(lech)} đơn vị BẤT ĐỒNG — cần hai người ngồi lại:")
        for i, x, y in lech:
            print(f"    · {i:<8} người 1 = {x:<6} người 2 = {y}")
    else:
        print("\n  Không có bất đồng nào.")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    phieu = Path(sys.argv[1])
    if not phieu.exists():
        print(f"Không tìm thấy {phieu}")
        sys.exit(1)

    cat = {r["id"]: r for r in json.loads(CAT.read_text(encoding="utf-8"))["rules"]}
    wb = openpyxl.load_workbook(phieu)
    ws = wb["Phiếu"] if "Phiếu" in wb.sheetnames else wb.active
    head = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(head)}

    cap_lop, ids_lop, cap_dl, ids_dl = [], [], [], []
    thieu = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rid = row[ix["id"]]
        if not rid:
            continue
        goc = cat.get(rid)
        if goc is None:
            print(f"  [cảnh báo] {rid} không có trong danh mục — bỏ qua")
            continue
        l2 = (row[ix["lop"]] or "").strip().upper() or None
        d2 = (row[ix["du_lieu"]] or "").strip().lower() or None
        if l2 is None and d2 is None:
            thieu.append(rid)
            continue
        if l2:
            cap_lop.append((goc["lop"], l2)); ids_lop.append(rid)
        if d2:
            cap_dl.append((goc["du_lieu"], d2)); ids_dl.append(rid)

    print("=" * 74)
    print("MỨC ĐỒNG THUẬN GIỮA HAI NGƯỜI GÁN NHÃN — R-Formal (CPA² Tuần 5)")
    print(f"Phiếu: {phieu.name}")
    print("=" * 74)
    if thieu:
        print(f"  {len(thieu)} đơn vị CHƯA ĐIỀN: {', '.join(thieu[:12])}"
              + (" ..." if len(thieu) > 12 else ""))

    bao_cao("Khả năng hình thức hoá (F / P / N)", cap_lop, ["F", "P", "N"], ids_lop)
    bao_cao("Khả dụng dữ kiện (co / khong)", cap_dl, ["co", "khong"], ids_dl)

    print("\n" + "=" * 74)
    print("CÁCH DÙNG KẾT QUẢ TRONG LUẬN VĂN")
    print("=" * 74)
    print("  · Báo cáo kappa cùng với cỡ mẫu và cách chọn mẫu — kappa không kèm n là vô nghĩa.")
    print("  · kappa < 0,60 nghĩa là tiêu chí phân loại chưa đủ rõ: sửa ĐỊNH NGHĨA trong")
    print("    taxonomy rồi gán lại, đừng chỉ ghi nhận con số thấp.")
    print("  · Các đơn vị bất đồng thường là ca thú vị nhất — chúng thường lộ ra một loại")
    print("    lý do chưa có trong khung phân loại. Nên ghi lại nội dung tranh luận.")
    print("  · Nhãn chốt sau khi thống nhất mới là nhãn đưa vào r_formal_v1.json.")


if __name__ == "__main__":
    main()
